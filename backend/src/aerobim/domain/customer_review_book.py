"""Ingest customer RD review books by sheet structure. Not gold. No auto-match.

PII (FIO, email) is stripped on the way in. Output is requirement rows plus
round observations, never a finding↔remark matrix. Checkpoint GO; customer_go false.
"""

from __future__ import annotations

import hashlib
import re
from collections.abc import Mapping, Sequence
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Final, Literal

from aerobim.domain.checkpoint import CHECKPOINT, CUSTOMER_GO

CLAIM_BOUNDARY: Final = (
    "RD acceptance-checklist rows. One adjudicator. Not gold. "
    "Not product accuracy. No automatic finding match. "
    "Checkpoint GO; customer_go false."
)

SourceKind = Literal["review_book", "requirement_template"]
EMAIL_RE: Final = re.compile(r"[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}", re.IGNORECASE)
PII_HEADER_RE: Final = re.compile(
    r"(фио|проверяющ|проектировщ|e-?mail|почта|должность|специалист)",
    re.IGNORECASE,
)
STATUS_HEADER: Final = "статус проверки"
REQUIREMENT_HEADER: Final = "содержание требований"
COMMENT_HEADER: Final = "комментарии"
ROUND_MARKERS: Final[tuple[tuple[str, int], ...]] = (
    ("первичн", 0),
    ("повторн", 1),
    ("критическ", 2),
)
SAMPLES_CUSTOMER: Final = "samples/customer"


@dataclass(frozen=True)
class RequirementObservation:
    round_index: int
    status: str
    comment: str


@dataclass(frozen=True)
class RequirementRow:
    row_id: str
    source_kind: SourceKind
    book_sha: str
    kit_code: str
    checklist_section: str
    requirement_no: str
    requirement_text: str
    sheet_ref: str
    stage: str
    stage_mismatch: bool
    observations: tuple[RequirementObservation, ...]
    matched_finding_id: str = ""


class ReviewBookError(ValueError):
    """Unknown header or illegal output path."""


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def dedupe_books(paths: Sequence[Path]) -> tuple[tuple[Path, ...], int]:
    seen: dict[str, Path] = {}
    duplicates = 0
    for path in paths:
        digest = sha256_file(path)
        if digest in seen:
            duplicates += 1
            continue
        seen[digest] = path
    return tuple(seen.values()), duplicates


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.upper() in {"#NAME?", "#REF!", "#VALUE!", "NONE"}:
        return ""
    return EMAIL_RE.sub("", text).strip()


def _is_pii_header(text: str) -> bool:
    return bool(PII_HEADER_RE.search(text or ""))


def _norm(text: str) -> str:
    return " ".join((text or "").casefold().split())


def classify_sheet_kind(headers: Sequence[str], status_values: Sequence[str]) -> SourceKind | None:
    joined = " | ".join(_norm(item) for item in headers if item)
    has_requirement = REQUIREMENT_HEADER in joined
    has_status = STATUS_HEADER in joined
    filled_status = any(item.strip() for item in status_values)
    if has_status and filled_status:
        return "review_book"
    if has_requirement and not filled_status:
        return "requirement_template"
    return None


def _load_sheet_rows(path: Path) -> list[tuple[str, list[list[str]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: list[tuple[str, list[list[str]]]] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for raw in sheet.iter_rows(values_only=True):
                rows.append([_cell_text(cell) for cell in raw])
            sheets.append((sheet.title, rows))
    finally:
        workbook.close()
    return sheets


def _header_index(rows: Sequence[Sequence[str]]) -> int | None:
    for index, row in enumerate(rows[:20]):
        joined = " ".join(_norm(cell) for cell in row)
        if STATUS_HEADER in joined or REQUIREMENT_HEADER in joined:
            return index
    return None


def _column_roles(header_row: Sequence[str], round_row: Sequence[str]) -> dict[str, Any]:
    req_col: int | None = None
    no_col: int | None = None
    status_cols: list[tuple[int, int]] = []
    comment_cols: dict[int, int] = {}
    current_round = 0
    for index, cell in enumerate(header_row):
        if _is_pii_header(cell):
            continue
        label = _norm(cell)
        round_label = _norm(round_row[index]) if index < len(round_row) else ""
        for marker, round_index in ROUND_MARKERS:
            if marker in round_label:
                current_round = round_index
                break
        if REQUIREMENT_HEADER in label or "лист рд" in label:
            if req_col is None:
                req_col = index
            continue
        if label in {"№№", "№", "no", "n"}:
            no_col = index
            continue
        if STATUS_HEADER in label:
            status_cols.append((index, current_round))
            continue
        if COMMENT_HEADER in label or "замечан" in label:
            comment_cols[current_round] = index
    return {
        "requirement": req_col,
        "number": no_col,
        "status_cols": status_cols,
        "comment_cols": comment_cols,
        "headers": list(header_row),
    }


def parse_review_book(
    path: Path,
    *,
    rd_ifc_present: bool = False,
    coverage_hash: str | None = None,
    expected_coverage_hash: str | None = None,
    column_map: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    if expected_coverage_hash and coverage_hash and coverage_hash != expected_coverage_hash:
        raise ReviewBookError("coverage map hash changed after freeze; ingest refused")
    book_sha = sha256_file(path)
    sheets = _load_sheet_rows(path)
    rows_out: list[RequirementRow] = []
    source_kind: SourceKind | None = None
    for sheet_name, grid in sheets:
        if not grid:
            continue
        header_i = _header_index(grid)
        if header_i is None:
            headers_found = [cell for row in grid[:8] for cell in row if cell]
            raise ReviewBookError(f"unknown header; found: {headers_found[:20]}")
        header_row = grid[header_i]
        round_row = grid[header_i - 1] if header_i > 0 else []
        roles = dict(column_map) if column_map else _column_roles(header_row, round_row)
        req_col = roles.get("requirement")
        if req_col is None:
            raise ReviewBookError(f"unknown header; found: {header_row}")
        status_cols: list[tuple[int, int]] = list(roles.get("status_cols") or ())
        comment_cols: dict[int, int] = dict(roles.get("comment_cols") or {})
        status_values: list[str] = []
        for data in grid[header_i + 1 :]:
            for col, _round in status_cols:
                if col < len(data):
                    status_values.append(data[col])
        kind = classify_sheet_kind(header_row, status_values)
        if kind is None:
            continue
        source_kind = kind if source_kind is None else source_kind
        section = ""
        for data in grid[header_i + 1 :]:
            req_text = data[req_col] if req_col < len(data) else ""
            if not req_text:
                joined = " ".join(data)
                if "замечаний в комплекте" in _norm(joined):
                    continue
                if (
                    req_text == ""
                    and any(data)
                    and not any(data[col] if col < len(data) else "" for col, _ in status_cols)
                ):
                    # lettered checklist section heading
                    section = next((cell for cell in data if cell), section)
                continue
            number = ""
            no_col = roles.get("number")
            if isinstance(no_col, int) and no_col < len(data):
                number = data[no_col]
            observations: list[RequirementObservation] = []
            if status_cols:
                for col, round_index in status_cols:
                    status = data[col] if col < len(data) else ""
                    comment_i = comment_cols.get(round_index)
                    comment = (
                        data[comment_i] if comment_i is not None and comment_i < len(data) else ""
                    )
                    observations.append(
                        RequirementObservation(
                            round_index=round_index,
                            status=status,
                            comment=comment,
                        )
                    )
            else:
                observations.append(RequirementObservation(round_index=0, status="", comment=""))
            stage = (
                "РД" if "рд" in _norm(req_text) or "лист рд" in _norm(" ".join(header_row)) else ""
            )
            stage_mismatch = (stage == "РД" or "рд" in _norm(sheet_name)) and not rd_ifc_present
            row_id = hashlib.sha256(
                f"{book_sha}:{sheet_name}:{number}:{req_text}".encode()
            ).hexdigest()[:16]
            rows_out.append(
                RequirementRow(
                    row_id=row_id,
                    source_kind=kind,
                    book_sha=book_sha,
                    kit_code="",
                    checklist_section=section,
                    requirement_no=number,
                    requirement_text=req_text,
                    sheet_ref="",
                    stage=stage or "РД",
                    stage_mismatch=stage_mismatch,
                    observations=tuple(observations),
                    matched_finding_id="",
                )
            )
    if source_kind is None:
        raise ReviewBookError("unknown header; no requirement/status columns")
    payload_rows = []
    for item in rows_out:
        payload = asdict(item)
        payload["matched_finding_id"] = ""
        payload_rows.append(payload)
    return {
        "schema_version": "1.0.0",
        "artifact_type": "customer_review_book",
        "source_kind": source_kind,
        "book_sha": book_sha,
        "claim_boundary": CLAIM_BOUNDARY,
        "checkpoint": CHECKPOINT,
        "customer_go": CUSTOMER_GO,
        "adjudicators": 1,
        "publishable": False,
        "is_accuracy": False,
        "requirement_count": len(payload_rows),
        "rows": payload_rows,
    }


def assert_output_path_allowed(output_dir: Path, *, repo: Path) -> None:
    resolved = output_dir.resolve()
    customer = (repo / SAMPLES_CUSTOMER).resolve()
    try:
        resolved.relative_to(customer)
    except ValueError:
        return
    raise ReviewBookError("refuses to write into samples/customer")


def strip_pii_blob(text: str) -> str:
    return EMAIL_RE.sub("", text)


def eligible_for_fn(*, coverage_status: str, stage_mismatch: bool) -> bool:
    if stage_mismatch:
        return False
    return coverage_status == "checked"


def observation_key(row: Mapping[str, Any]) -> str:
    text = str(row.get("requirement_text") or "")
    comments = []
    for obs in row.get("observations") or ():
        if isinstance(obs, Mapping):
            comments.append(str(obs.get("comment") or ""))
    return f"{text}|{'|'.join(comments)}"
